import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from agents.export_agent import ExportAgent
from schemas.company import Company


def _company(**overrides) -> Company:
    defaults = dict(
        company_name="Focet Valves Private Limited",
        gst="27FOCET1234F1Z5",
        website="https://focetvalves.example",
        phone="9876543210",
        email="sales@focetvalves.example",
        city="Coimbatore",
    )
    defaults.update(overrides)
    return Company(**defaults)


def _write_template(path: Path, headers: list[str], sheet_name: str = "Leads") -> None:
    wb = Workbook()
    sheet = wb.active
    sheet.title = sheet_name
    sheet.append(headers)
    wb.save(path)


class ExportAgentTemplateTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def _export(self, companies, template_headers, output_name="output.xlsx", sheet_name="Leads"):
        template_path = self.tmp_path / "template.xlsx"
        _write_template(template_path, template_headers, sheet_name=sheet_name)
        output_path = self.tmp_path / output_name
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(template_path)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute(companies, output_path=str(output_path))
        return output_path

    # TEST 1: default/current template -> correct output columns and values.
    def test_default_template_columns_and_values(self):
        headers = ["Company Name", "GST", "Website URL", "Mobile Number", "Email ID", "City"]
        out = self._export([_company()], headers)
        wb = load_workbook(out)
        sheet = wb["Leads"]
        self.assertEqual([c.value for c in sheet[1]], headers)
        row = [c.value for c in sheet[2]]
        self.assertEqual(row, [
            "Focet Valves Private Limited", "27FOCET1234F1Z5",
            "https://focetvalves.example", "9876543210",
            "sales@focetvalves.example", "Coimbatore",
        ])

    # TEST 2: reordered columns -> output follows exactly that order.
    def test_reordered_columns_are_respected(self):
        headers = ["Email", "Company", "GST", "Website"]
        out = self._export([_company()], headers)
        wb = load_workbook(out)
        sheet = wb["Leads"]
        self.assertEqual([c.value for c in sheet[1]], headers)
        row = [c.value for c in sheet[2]]
        self.assertEqual(row, [
            "sales@focetvalves.example", "Focet Valves Private Limited",
            "27FOCET1234F1Z5", "https://focetvalves.example",
        ])

    # TEST 3: renamed headers -> correct field mapping via alias registry.
    def test_renamed_headers_map_correctly(self):
        headers = ["Company", "GST Number", "Website URL", "Contact Number"]
        out = self._export([_company()], headers)
        wb = load_workbook(out)
        sheet = wb["Leads"]
        row = [c.value for c in sheet[2]]
        self.assertEqual(row, [
            "Focet Valves Private Limited", "27FOCET1234F1Z5",
            "https://focetvalves.example", "9876543210",
        ])

    # TEST 4: template has an extra column with no extracted field -> blank, no crash.
    def test_unmapped_template_column_left_blank(self):
        headers = ["Company Name", "GST", "Annual Revenue"]
        out = self._export([_company()], headers)
        wb = load_workbook(out)
        sheet = wb["Leads"]
        row = [c.value for c in sheet[2]]
        self.assertEqual(row, ["Focet Valves Private Limited", "27FOCET1234F1Z5", None])

    # TEST 5: extracted data has a field not present in the template -> ignored,
    # no unwanted column is added.
    def test_extra_extracted_field_is_ignored(self):
        headers = ["Company Name", "GST"]
        company = _company(remarks="internal note", followup="call next week")
        out = self._export([company], headers)
        wb = load_workbook(out)
        sheet = wb["Leads"]
        self.assertEqual([c.value for c in sheet[1]], headers)
        self.assertEqual(sheet.max_column, 2)

    # TEST 6: replacing the template with a different structure requires no
    # Python code change - the second export follows the new template, and
    # the mismatched old output is archived rather than corrupted/dropped.
    def test_template_replacement_requires_no_code_change(self):
        output_path = self.tmp_path / "master.xlsx"

        template_a = self.tmp_path / "template_a.xlsx"
        _write_template(template_a, ["Company Name", "GST"])
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(template_a)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute([_company()], output_path=str(output_path))

        first_wb = load_workbook(output_path)
        self.assertEqual([c.value for c in first_wb["Leads"][1]], ["Company Name", "GST"])

        # Business swaps in a structurally different template - no code change.
        template_b = self.tmp_path / "template_b.xlsx"
        _write_template(template_b, ["Company", "Industry Type", "Region", "Contact Person"])
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(template_b)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute(
                [_company(industry="Valves", region="South", contact_person="A. Kumar")],
                output_path=str(output_path),
            )

        second_wb = load_workbook(output_path)
        self.assertEqual(
            [c.value for c in second_wb["Leads"][1]],
            ["Company", "Industry Type", "Region", "Contact Person"],
        )
        self.assertEqual(
            [c.value for c in second_wb["Leads"][2]],
            ["Focet Valves Private Limited", "Valves", "South", "A. Kumar"],
        )

        backups = list(self.tmp_path.glob("master.backup_*.xlsx"))
        self.assertEqual(len(backups), 1, "mismatched previous output should be archived, not lost")
        archived_wb = load_workbook(backups[0])
        self.assertEqual([c.value for c in archived_wb["Leads"][1]], ["Company Name", "GST"])


class ProductionLeadCommonImportDTOTemplateTests(unittest.TestCase):
    """Exercises the actual 16-column production template (sheet
    "LeadCommonImportDTO") bundled at backend/templates/LeadImport_template.xlsx,
    including its embedded hidden mapping sheet: "Company" is a
    template-defined STATIC value ("Goodbooks"), "Lead Name" carries the
    extracted company_name, "Turnover" carries Company.turnover, and
    Description/Stage/Incharge/Alloted To/Source Detail are Marketing-owned
    columns that must remain unmapped rather than fabricated."""

    TEMPLATE_HEADERS = [
        "Company", "Lead Name", "Description", "Contact Name", "Designation",
        "Contact Mobile No", "Contact Email", "Address1", "City", "GST",
        "Turnover", "Remarks", "Stage", "Incharge", "Alloted To", "Source Detail",
    ]

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        backend_dir = Path(__file__).resolve().parent.parent
        self.template_path = backend_dir / "templates" / "LeadImport_template.xlsx"
        self.assertTrue(self.template_path.exists(), "bundled production template is missing")

    def _export(self, companies):
        output_path = self.tmp_path / "output.xlsx"
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(self.template_path)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute(companies, output_path=str(output_path))
        return output_path

    def test_column_order_matches_uploaded_template_exactly(self):
        out = self._export([_company()])
        wb = load_workbook(out)
        sheet = wb["LeadCommonImportDTO"]
        self.assertEqual([c.value for c in sheet[1]], self.TEMPLATE_HEADERS)

    def test_known_fields_map_correctly(self):
        company = _company(
            contact_person="A. Kumar",
            designation="Purchase Manager",
            address="12 Industrial Estate",
            remarks="Interested in bulk order",
            turnover="25 Cr",
        )
        out = self._export([company])
        wb = load_workbook(out)
        sheet = wb["LeadCommonImportDTO"]
        headers = [c.value for c in sheet[1]]
        row = dict(zip(headers, [c.value for c in sheet[2]]))
        # "Company" is a template-defined STATIC value, never the
        # extracted company name.
        self.assertEqual(row["Company"], "Goodbooks")
        # "Lead Name" is where the extracted company name actually lands.
        self.assertEqual(row["Lead Name"], "Focet Valves Private Limited")
        self.assertEqual(row["Contact Name"], "A. Kumar")
        self.assertEqual(row["Designation"], "Purchase Manager")
        self.assertEqual(row["Contact Mobile No"], "9876543210")
        self.assertEqual(row["Contact Email"], "sales@focetvalves.example")
        self.assertEqual(row["GST"], "27FOCET1234F1Z5")
        self.assertEqual(row["City"], "Coimbatore")
        self.assertEqual(row["Turnover"], "25 Cr")
        self.assertEqual(row["Remarks"], "Interested in bulk order")
        self.assertEqual(row["Address1"], "12 Industrial Estate")

    def test_unsupported_fields_stay_blank_not_fabricated(self):
        out = self._export([_company()])
        wb = load_workbook(out)
        sheet = wb["LeadCommonImportDTO"]
        headers = [c.value for c in sheet[1]]
        row = dict(zip(headers, [c.value for c in sheet[2]]))
        for unmapped in ["Description", "Stage", "Incharge", "Alloted To", "Source Detail"]:
            self.assertIsNone(row[unmapped], f"{unmapped} must stay blank, not be fabricated")

    def test_bundled_template_sample_row_is_not_leaked_into_output(self):
        # The uploaded production template ships with an illustrative
        # "Goodbooks" example row (under the header, in the Company
        # column) - that row itself must never end up as a SECOND row of
        # generated output. "Goodbooks" legitimately appearing once, in
        # the Company column of the one real exported row, is correct -
        # it's the static value that column always carries.
        out = self._export([_company()])
        wb = load_workbook(out)
        sheet = wb["LeadCommonImportDTO"]
        self.assertEqual(sheet.max_row, 2, "only header + the one real company row should be present")
        self.assertEqual(sheet.cell(row=2, column=1).value, "Goodbooks")


class ExportAgentMasterFileTransitionTests(unittest.TestCase):
    """Reproduces the exact production scenario reported after swapping in
    backend/templates/LeadImport_template.xlsx: an existing master export
    file (from before the swap, still shaped like the OLD template - a
    'Leads' sheet with headers such as 'Alternate Mobile Number' and no
    'LeadCommonImportDTO' sheet at all) gets exported against with the
    CURRENT production template. Uses the real bundled template, not a
    synthetic one, so this fails if that specific transition ever breaks.
    """

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        backend_dir = Path(__file__).resolve().parent.parent
        self.template_path = backend_dir / "templates" / "LeadImport_template.xlsx"
        self.assertTrue(self.template_path.exists(), "bundled production template is missing")

    def _old_format_master(self) -> Path:
        path = self.tmp_path / "All_Extracted_Leads.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.title = "Leads"
        sheet.append([
            "Company Name", "GST", "Turn Over", "Region", "City", "Industry Type",
            "Contact Person", "Designation", "Mobile Number", "Alternate Mobile Number",
            "Email ID", "LinkedIN Id", "Website URL", "Remarks", "Followup", "Extracted On",
        ])
        sheet.append([
            "Old Format Co", "27OLD1234F1Z5", "5 Cr", "West", "Pune", "Valve",
            "A B", "Manager", "9990001111", "8880002222",
            "old@example.com", "", "", "", "", "2026-08-21 10:00",
        ])
        wb.save(path)
        return path

    def test_export_against_old_format_master_succeeds_and_archives_it(self):
        companies = [_company(company_name=f"New Valve Co {i}", phone_alt="8888888888") for i in range(10)]
        master_path = self._old_format_master()

        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(self.template_path)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            result_path = ExportAgent().execute(companies, output_path=str(master_path))

        # No exception, and the old-format file was archived rather than
        # silently overwritten or lost.
        backups = list(self.tmp_path.glob("All_Extracted_Leads.backup_*.xlsx"))
        self.assertEqual(len(backups), 1, "old-format master should be archived, not discarded")
        archived = load_workbook(backups[0])
        self.assertIn("Old Format Co", {c.value for row in archived["Leads"].iter_rows() for c in row})

        # The new output is built from the CURRENT template's structure.
        wb = load_workbook(result_path)
        sheet = wb["LeadCommonImportDTO"]
        headers = [c.value for c in sheet[1]]
        self.assertEqual(
            headers,
            [
                "Company", "Lead Name", "Description", "Contact Name", "Designation",
                "Contact Mobile No", "Contact Email", "Address1", "City", "GST",
                "Turnover", "Remarks", "Stage", "Incharge", "Alloted To", "Source Detail",
            ],
        )
        self.assertNotIn("Alternate Mobile Number", headers)
        self.assertEqual(sheet.max_row, 11, "header + 10 new companies, no leftover old rows")

        # phone_alt has no matching column in the new template and must be
        # silently ignored for export - not an error, and not fabricated
        # into some other column.
        row_values = {c.value for row in sheet.iter_rows(min_row=2) for c in row}
        self.assertNotIn("8888888888", row_values)


class ExportAgentLockedFileTests(unittest.TestCase):
    """A file locked/open in another program (Excel, OneDrive sync, an AV
    scan) is the most likely real-world cause of an OSError from
    Path.rename()/Workbook.save() during export. Previously that OSError
    propagated with no server-side logging at all - this proves it now
    surfaces as a clear, actionable RuntimeError instead."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.template_path = self.tmp_path / "template.xlsx"
        _write_template(self.template_path, ["Company", "GST"], sheet_name="Leads")

    def test_locked_output_file_raises_actionable_error(self):
        output_path = self.tmp_path / "output.xlsx"
        # Deliberately different headers than the template so the
        # archive-and-recreate branch (the one that calls Path.rename) is
        # the one under test.
        _write_template(output_path, ["Old Header A", "Old Header B"], sheet_name="Leads")

        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(self.template_path)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None), \
             patch("pathlib.Path.rename", side_effect=PermissionError("simulated: file is open elsewhere")):
            with self.assertRaises(RuntimeError) as ctx:
                ExportAgent().execute([_company()], output_path=str(output_path))
        self.assertIn("locked", str(ctx.exception).lower())


if __name__ == "__main__":
    unittest.main()


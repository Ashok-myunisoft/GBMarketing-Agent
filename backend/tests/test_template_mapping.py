"""Tests for services/excel_template/template_mapping.py and its
integration with template_reader.py / agents/export_agent.py.

These prove the actual business requirement: a template's own hidden
mapping sheet is enough to change what a column means (a real field, a
fixed value, or intentionally blank), and none of that requires touching
any .py file - renaming, reordering, adding, or removing template
columns, or swapping the entire template for a structurally different
one, only ever requires editing/replacing the .xlsx file.
"""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from agents.export_agent import ExportAgent
from schemas.company import Company
from services.excel_template.template_mapping import MAPPING_SHEET_NAME


def _company(**overrides) -> Company:
    defaults = dict(
        company_name="Focet Valves Private Limited",
        contact_person="A. Kumar",
        designation="Purchase Manager",
        phone="9876543210",
        email="sales@focetvalves.example",
        address="12 Industrial Estate",
        city="Coimbatore",
        gst="27FOCET1234F1Z5",
        turnover="25 Cr",
        remarks="Interested in bulk order",
    )
    defaults.update(overrides)
    return Company(**defaults)


def _build_template(path: Path, headers: list[str], mapping_rows: list[tuple] | None,
                     sample_row: list | None = None, sheet_name: str = "Sheet1") -> None:
    """mapping_rows: list of (template_column, maps_to_field, static_value)
    tuples. Pass None to build a template with NO mapping sheet at all
    (legacy alias-matching path)."""
    wb = Workbook()
    data_sheet = wb.active
    data_sheet.title = sheet_name
    data_sheet.append(headers)
    if sample_row is not None:
        data_sheet.append(sample_row)

    if mapping_rows is not None:
        mapping_sheet = wb.create_sheet(MAPPING_SHEET_NAME)
        mapping_sheet.append(["Template Column", "Maps To Field", "Static Value", "Notes"])
        for row in mapping_rows:
            mapping_sheet.append(list(row))
        mapping_sheet.sheet_state = "hidden"
        wb.active = wb.sheetnames.index(sheet_name)

    wb.save(path)


class TemplateMappingReadTests(unittest.TestCase):
    """Unit tests directly on template_mapping.read_mapping_overrides()."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_no_mapping_sheet_returns_none(self):
        from openpyxl import load_workbook
        from services.excel_template.template_mapping import read_mapping_overrides

        path = self.tmp_path / "plain.xlsx"
        _build_template(path, ["Company Name", "GST"], mapping_rows=None)
        wb = load_workbook(path)
        self.assertIsNone(read_mapping_overrides(wb))

    def test_static_field_and_blank_rows_are_classified_correctly(self):
        from openpyxl import load_workbook
        from services.excel_template.template_mapping import read_mapping_overrides

        path = self.tmp_path / "mapped.xlsx"
        _build_template(
            path,
            ["Company", "Lead Name", "Stage"],
            mapping_rows=[
                ("Company", "static", "Goodbooks"),
                ("Lead Name", "company_name", ""),
                ("Stage", "", ""),
            ],
        )
        wb = load_workbook(path)
        overrides = read_mapping_overrides(wb)
        self.assertIsNotNone(overrides)

        from services.excel_template.field_aliases import normalize
        self.assertEqual(overrides[normalize("Company")].kind, "static")
        self.assertEqual(overrides[normalize("Company")].static_value, "Goodbooks")
        self.assertEqual(overrides[normalize("Lead Name")].kind, "field")
        self.assertEqual(overrides[normalize("Lead Name")].field, "company_name")
        self.assertEqual(overrides[normalize("Stage")].kind, "blank")

    def test_invalid_field_reference_is_marked_invalid_not_raised(self):
        from openpyxl import load_workbook
        from services.excel_template.template_mapping import read_mapping_overrides

        path = self.tmp_path / "typo.xlsx"
        _build_template(
            path,
            ["GST"],
            mapping_rows=[("GST", "gst_number_typo", "")],  # not a real Company field
        )
        wb = load_workbook(path)
        overrides = read_mapping_overrides(wb)

        from services.excel_template.field_aliases import normalize
        self.assertEqual(overrides[normalize("GST")].kind, "invalid")


class TemplateMappingIntegrationTests(unittest.TestCase):
    """Full read_template_layout() + ExportAgent integration, proving the
    business requirements end-to-end without touching any .py file."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def _export(self, headers, mapping_rows, companies, sample_row=None):
        from unittest.mock import patch

        template_path = self.tmp_path / "template.xlsx"
        _build_template(template_path, headers, mapping_rows, sample_row=sample_row)
        output_path = self.tmp_path / "output.xlsx"
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(template_path)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute(companies, output_path=str(output_path))
        return output_path

    def _read_row(self, output_path, row_number=2):
        from openpyxl import load_workbook
        wb = load_workbook(output_path)
        sheet = wb.active
        headers = [c.value for c in sheet[1]]
        return headers, dict(zip(headers, [c.value for c in sheet[row_number]]))

    # --- Requirement: static values -----------------------------------
    def test_static_value_is_written_verbatim_for_every_row(self):
        headers = ["Company", "Lead Name"]
        mapping = [("Company", "static", "Goodbooks"), ("Lead Name", "company_name", "")]
        out = self._export(headers, mapping, [_company(company_name="Alpha Co"), _company(company_name="Beta Co")])
        from openpyxl import load_workbook
        sheet = load_workbook(out).active
        self.assertEqual([sheet.cell(row=r, column=1).value for r in (2, 3)], ["Goodbooks", "Goodbooks"])
        self.assertEqual([sheet.cell(row=r, column=2).value for r in (2, 3)], ["Alpha Co", "Beta Co"])

    def test_static_value_is_not_treated_as_an_extracted_lead(self):
        # The "Goodbooks" sample row under the header must never itself
        # become a row of exported lead data.
        headers = ["Company", "Lead Name"]
        mapping = [("Company", "static", "Goodbooks"), ("Lead Name", "company_name", "")]
        out = self._export(headers, mapping, [_company(company_name="Real Co")], sample_row=["Goodbooks", None])
        from openpyxl import load_workbook
        sheet = load_workbook(out).active
        self.assertEqual(sheet.max_row, 2, "only header + the one real company row")
        self.assertEqual(sheet.cell(row=2, column=2).value, "Real Co")

    # --- Requirement: Lead Name / Turnover mappings ---------------------
    def test_lead_name_receives_company_name(self):
        headers = ["Lead Name"]
        mapping = [("Lead Name", "company_name", "")]
        _, row = self._read_row(self._export(headers, mapping, [_company(company_name="Zenith Valves")]))
        self.assertEqual(row["Lead Name"], "Zenith Valves")

    def test_turnover_receives_turnover(self):
        headers = ["Turnover"]
        mapping = [("Turnover", "turnover", "")]
        _, row = self._read_row(self._export(headers, mapping, [_company(turnover="42 Cr")]))
        self.assertEqual(row["Turnover"], "42 Cr")

    # --- Requirement: marketing-owned columns ---------------------------
    def test_marketing_owned_columns_stay_blank(self):
        headers = ["Lead Name", "Stage", "Incharge"]
        mapping = [("Lead Name", "company_name", ""), ("Stage", "", ""), ("Incharge", "", "")]
        _, row = self._read_row(self._export(headers, mapping, [_company()]))
        self.assertIsNone(row["Stage"])
        self.assertIsNone(row["Incharge"])

    # --- Requirement: ambiguous/unmapped columns ------------------------
    def test_ambiguous_field_reference_stays_blank_not_fabricated(self):
        headers = ["Lead Name", "Weird Column"]
        mapping = [("Lead Name", "company_name", ""), ("Weird Column", "not_a_real_field", "")]
        _, row = self._read_row(self._export(headers, mapping, [_company(company_name="X Co")]))
        self.assertEqual(row["Lead Name"], "X Co")
        self.assertIsNone(row["Weird Column"])

    def test_column_missing_from_mapping_sheet_stays_blank(self):
        # "Extra Column" is physically present in the data sheet but has
        # no row in the mapping sheet at all - must NOT silently fall
        # back to alias-guessing once a template opts into explicit
        # mapping.
        headers = ["Lead Name", "Extra Column"]
        mapping = [("Lead Name", "company_name", "")]  # "Extra Column" omitted on purpose
        _, row = self._read_row(self._export(headers, mapping, [_company(company_name="X Co")]))
        self.assertIsNone(row["Extra Column"])

    # --- Requirement: reordering/renaming/adding/removing columns ------
    def test_reordering_columns_requires_no_code_change(self):
        headers = ["Turnover", "Lead Name", "GST"]  # deliberately reordered
        mapping = [("Turnover", "turnover", ""), ("Lead Name", "company_name", ""), ("GST", "gst", "")]
        out_headers, row = self._read_row(
            self._export(headers, mapping, [_company(company_name="Reorder Co", turnover="9 Cr", gst="27ABC")])
        )
        self.assertEqual(out_headers, ["Turnover", "Lead Name", "GST"])
        self.assertEqual(row["Turnover"], "9 Cr")
        self.assertEqual(row["Lead Name"], "Reorder Co")
        self.assertEqual(row["GST"], "27ABC")

    def test_renaming_headers_requires_no_code_change(self):
        # Template B from the requirements: "Customer" / "Business Name" /
        # "Email Address" / "GST Number" / "Revenue" - all custom header
        # text, entirely mapping-sheet-driven.
        headers = ["Customer", "Business Name", "Email Address", "GST Number", "Revenue"]
        mapping = [
            ("Customer", "static", "Goodbooks"),
            ("Business Name", "company_name", ""),
            ("Email Address", "email", ""),
            ("GST Number", "gst", ""),
            ("Revenue", "turnover", ""),
        ]
        _, row = self._read_row(
            self._export(headers, mapping, [_company(company_name="Renamed Co", email="a@b.com",
                                                       gst="27XYZ", turnover="15 Cr")])
        )
        self.assertEqual(row["Customer"], "Goodbooks")
        self.assertEqual(row["Business Name"], "Renamed Co")
        self.assertEqual(row["Email Address"], "a@b.com")
        self.assertEqual(row["GST Number"], "27XYZ")
        self.assertEqual(row["Revenue"], "15 Cr")

    def test_adding_a_new_template_column_requires_no_code_change(self):
        headers = ["Lead Name", "LinkedIn"]  # LinkedIn is new, not in the old template
        mapping = [("Lead Name", "company_name", ""), ("LinkedIn", "linkedin_url", "")]
        _, row = self._read_row(
            self._export(headers, mapping, [_company(company_name="New Col Co",
                                                       linkedin_url="https://linkedin.example/x")])
        )
        self.assertEqual(row["LinkedIn"], "https://linkedin.example/x")

    def test_removing_a_template_column_requires_no_code_change(self):
        # A template that simply never mentions e.g. "Designation" at all.
        headers = ["Lead Name", "Contact Mobile No"]
        mapping = [("Lead Name", "company_name", ""), ("Contact Mobile No", "phone", "")]
        out_headers, row = self._read_row(
            self._export(headers, mapping, [_company(company_name="Slim Co", phone="9000000000")])
        )
        self.assertEqual(out_headers, ["Lead Name", "Contact Mobile No"])
        self.assertEqual(row["Contact Mobile No"], "9000000000")

    def test_third_template_shape_from_requirements(self):
        # Template C: "Lead Name | Company | Phone | City | Turnover |
        # Marketing Status" - a full third, independently-shaped example.
        headers = ["Lead Name", "Company", "Phone", "City", "Turnover", "Marketing Status"]
        mapping = [
            ("Lead Name", "company_name", ""),
            ("Company", "static", "Goodbooks"),
            ("Phone", "phone", ""),
            ("City", "city", ""),
            ("Turnover", "turnover", ""),
            ("Marketing Status", "", ""),
        ]
        out_headers, row = self._read_row(
            self._export(headers, mapping, [_company(company_name="Template C Co", phone="9111111111",
                                                       city="Chennai", turnover="8 Cr")])
        )
        self.assertEqual(out_headers, headers)
        self.assertEqual(row["Lead Name"], "Template C Co")
        self.assertEqual(row["Company"], "Goodbooks")
        self.assertEqual(row["Phone"], "9111111111")
        self.assertEqual(row["City"], "Chennai")
        self.assertEqual(row["Turnover"], "8 Cr")
        self.assertIsNone(row["Marketing Status"])

    # --- Requirement: full template replacement -------------------------
    def test_full_template_replacement_requires_no_code_change(self):
        from unittest.mock import patch

        template_a = self.tmp_path / "template_a.xlsx"
        _build_template(
            template_a, ["Company", "Lead Name", "Email", "GST"],
            mapping_rows=[
                ("Company", "static", "Goodbooks"),
                ("Lead Name", "company_name", ""),
                ("Email", "email", ""),
                ("GST", "gst", ""),
            ],
        )
        template_c = self.tmp_path / "template_c.xlsx"
        _build_template(
            template_c, ["Lead Name", "Company", "Phone", "City", "Turnover", "Marketing Status"],
            mapping_rows=[
                ("Lead Name", "company_name", ""),
                ("Company", "static", "Goodbooks"),
                ("Phone", "phone", ""),
                ("City", "city", ""),
                ("Turnover", "turnover", ""),
                ("Marketing Status", "", ""),
            ],
        )

        company = _company(company_name="Swap Co", email="swap@example.com", gst="27SWAP",
                            phone="9222222222", city="Mumbai", turnover="11 Cr")

        out_a = self.tmp_path / "out_a.xlsx"
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(template_a)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute([company], output_path=str(out_a))

        out_c = self.tmp_path / "out_c.xlsx"
        with patch("agents.export_agent.settings.EXCEL_TEMPLATE_PATH", str(template_c)), \
             patch("agents.export_agent.settings.EXCEL_TEMPLATE_HEADER_ROW", None):
            ExportAgent().execute([company], output_path=str(out_c))

        _, row_a = self._read_row(out_a)
        _, row_c = self._read_row(out_c)
        self.assertEqual(row_a["Lead Name"], "Swap Co")
        self.assertEqual(row_a["Email"], "swap@example.com")
        self.assertEqual(row_c["Lead Name"], "Swap Co")
        self.assertEqual(row_c["City"], "Mumbai")
        # Same ExportAgent.execute() call, same export_agent.py code,
        # two structurally different templates - no code path differs.


if __name__ == "__main__":
    unittest.main()

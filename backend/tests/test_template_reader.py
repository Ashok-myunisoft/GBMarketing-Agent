"""Tests for services/excel_template/template_reader.py.

Covers: the module imports and exposes the objects agents/export_agent.py
depends on (TemplateLayout, read_template_layout), the bundled production
template (backend/templates/LeadImport_template.xlsx) is read correctly
via its embedded hidden mapping sheet - static values, field mappings,
and intentionally-blank marketing-owned columns all resolve as that
mapping sheet declares - and, independently, that templates with NO
mapping sheet at all still fall back to the older field_aliases.py
alias-matching behaviour unchanged.
"""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


class TemplateReaderImportTests(unittest.TestCase):
    def test_module_imports_successfully(self):
        from services.excel_template.template_reader import TemplateLayout, read_template_layout
        self.assertTrue(callable(read_template_layout))
        self.assertTrue(TemplateLayout)


class ProductionTemplateReadTests(unittest.TestCase):
    """Exercises the actual bundled template (with its embedded hidden
    __lead_export_mapping__ sheet) so this test fails loudly if the real
    file or its mapping sheet ever changes shape unexpectedly."""

    EXPECTED_HEADERS = [
        "Company", "Lead Name", "Description", "Contact Name", "Designation",
        "Contact Mobile No", "Contact Email", "Address1", "City", "GST",
        "Turnover", "Remarks", "Stage", "Incharge", "Alloted To", "Source Detail",
    ]

    def setUp(self):
        backend_dir = Path(__file__).resolve().parent.parent
        self.template_path = backend_dir / "templates" / "LeadImport_template.xlsx"
        self.assertTrue(self.template_path.exists(), "bundled production template is missing")

    def _read(self):
        from services.excel_template.template_reader import read_template_layout
        # header_row=None mirrors settings.EXCEL_TEMPLATE_HEADER_ROW's
        # default (auto-detect) - the same call agents/export_agent.py makes.
        return read_template_layout(self.template_path, header_row=None)

    def test_template_reads_successfully(self):
        layout = self._read()
        self.assertEqual(layout.sheet_name, "LeadCommonImportDTO")

    def test_header_row_is_detected_correctly(self):
        # Row 1 is the real header; row 2 is a single illustrative sample
        # value ("Goodbooks") that must not be mistaken for the header.
        layout = self._read()
        self.assertEqual(layout.header_row, 1)
        self.assertEqual(layout.header_values, self.EXPECTED_HEADERS)

    def test_company_column_is_a_static_value_not_extracted(self):
        # The mapping sheet declares "Company" -> static "Goodbooks" -
        # this must NOT resolve to Company.company_name, unlike the old
        # alias-only behaviour.
        layout = self._read()
        by_header = {c.header: c for c in layout.columns}
        company_column = by_header["Company"]
        self.assertIsNone(company_column.field)
        self.assertEqual(company_column.static_value, "Goodbooks")

    def test_lead_name_maps_to_company_name(self):
        layout = self._read()
        by_header = {c.header: c for c in layout.columns}
        self.assertEqual(by_header["Lead Name"].field, "company_name")
        self.assertIsNone(by_header["Lead Name"].static_value)

    def test_turnover_maps_to_turnover(self):
        layout = self._read()
        by_header = {c.header: c for c in layout.columns}
        self.assertEqual(by_header["Turnover"].field, "turnover")

    def test_known_headers_map_to_expected_company_fields(self):
        layout = self._read()
        by_header = {column.header: column.field for column in layout.columns}
        self.assertEqual(by_header["Contact Name"], "contact_person")
        self.assertEqual(by_header["Designation"], "designation")
        self.assertEqual(by_header["Contact Mobile No"], "phone")
        self.assertEqual(by_header["Contact Email"], "email")
        self.assertEqual(by_header["Address1"], "address")
        self.assertEqual(by_header["GST"], "gst")
        self.assertEqual(by_header["City"], "city")
        self.assertEqual(by_header["Remarks"], "remarks")

    def test_marketing_owned_columns_are_blank_not_fabricated(self):
        layout = self._read()
        by_header = {column.header: column for column in layout.columns}
        for unmapped in ["Description", "Stage", "Incharge", "Alloted To", "Source Detail"]:
            self.assertIn(unmapped, by_header, f"{unmapped} must keep its position in the layout")
            self.assertIsNone(by_header[unmapped].field, f"{unmapped} must have no Company field")
            self.assertIsNone(by_header[unmapped].static_value, f"{unmapped} must have no static value")

    def test_column_indexes_preserve_template_order(self):
        layout = self._read()
        self.assertEqual([column.index for column in layout.columns], list(range(1, len(self.EXPECTED_HEADERS) + 1)))
        self.assertEqual([column.header for column in layout.columns], layout.header_values)

    def test_mapping_sheet_itself_is_hidden_and_not_treated_as_data(self):
        from openpyxl import load_workbook
        from services.excel_template.template_mapping import MAPPING_SHEET_NAME

        wb = load_workbook(self.template_path)
        self.assertIn(MAPPING_SHEET_NAME, wb.sheetnames)
        self.assertEqual(wb[MAPPING_SHEET_NAME].sheet_state, "hidden")

        layout = self._read()
        self.assertEqual(layout.sheet_name, "LeadCommonImportDTO")
        self.assertNotEqual(layout.sheet_name, MAPPING_SHEET_NAME)


class HeaderDetectionAndAliasTests(unittest.TestCase):
    """Uses small throwaway workbooks with NO mapping sheet, so header-row
    auto-detection and the legacy alias-fallback path are verified
    independently of the production file and its mapping sheet."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_explicit_header_row_is_respected(self):
        from services.excel_template.template_reader import read_template_layout

        path = self.tmp_path / "explicit.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.append(["A random title row, not headers"])
        sheet.append(["Company Name", "GST", "Email"])
        wb.save(path)

        layout = read_template_layout(path, header_row=2)
        self.assertEqual(layout.header_row, 2)
        self.assertEqual(layout.header_values, ["Company Name", "GST", "Email"])

    def test_header_row_auto_detected_over_sparse_rows(self):
        from services.excel_template.template_reader import read_template_layout

        path = self.tmp_path / "auto.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.append(["Sample Co"])  # sparse row above the real header
        sheet.append(["Company Name", "GST", "Website URL", "Mobile Number"])
        wb.save(path)

        layout = read_template_layout(path, header_row=None)
        self.assertEqual(layout.header_row, 2)
        self.assertEqual(
            layout.header_values,
            ["Company Name", "GST", "Website URL", "Mobile Number"],
        )

    def test_renamed_header_still_maps_via_alias_registry(self):
        from services.excel_template.template_reader import read_template_layout

        path = self.tmp_path / "aliases.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.append(["Business Name", "GSTIN", "Contact Mobile No", "Annual Revenue"])
        wb.save(path)

        layout = read_template_layout(path, header_row=None)
        by_header = {column.header: column.field for column in layout.columns}
        self.assertEqual(by_header["Business Name"], "company_name")
        self.assertEqual(by_header["GSTIN"], "gst")
        self.assertEqual(by_header["Contact Mobile No"], "phone")
        # "Annual Revenue" has no registered alias - must stay unmapped,
        # not be guessed at, and must still keep its column position.
        self.assertIsNone(by_header["Annual Revenue"])


if __name__ == "__main__":
    unittest.main()
